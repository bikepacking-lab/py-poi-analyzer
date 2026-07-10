#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from pathlib import Path
from typing import Dict, Any, Optional
import logging
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core import gpx_processor
from ..base_generator import BaseGenerator

logger = logging.getLogger(__name__)

class ExcelGenerator(BaseGenerator):
    
    def get_format_type(self) -> str:
        return 'xlsx'
    
    def generate(self, data: Dict[str, Any], output_dir: Path, **kwargs) -> Optional[Path]:
        output_file = output_dir / f'analysis_report_{data.get("analysis_id", "unknown")}.xlsx'
        
        try:
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                self._create_summary_sheet(writer, data)
                self._create_poi_sheet(writer, data)
                self._create_stats_sheet(writer, data)
                self._create_config_sheet(writer, data)
            
            self._apply_styles(output_file)
            
            logger.info(f"Rapport Excel genere: {output_file}")
            return output_file
        except Exception as e:
            logger.error(f"Erreur Excel: {e}")
            return None
    
    def _get_value(self, obj: Any, key: str, default: str = '') -> str:
        """Recupere une valeur string de maniere sure"""
        if obj is None:
            return default
        if isinstance(obj, dict):
            val = obj.get(key, default)
            return str(val) if val is not None else default
        if hasattr(obj, key):
            val = getattr(obj, key, default)
            return str(val) if val is not None else default
        return default
    
    def _get_float(self, obj: Any, key: str, default: float = 0.0) -> float:
        """Recupere une valeur float de maniere sure"""
        if obj is None:
            return default
        try:
            if isinstance(obj, dict):
                val = obj.get(key, default)
            elif hasattr(obj, key):
                val = getattr(obj, key, default)
            else:
                return default
            if val is None:
                return default
            return float(val)
        except (ValueError, TypeError):
            return default
    
    def _create_summary_sheet(self, writer, data):
        """Cree la feuille de resume avec les temps"""
        try:
            summary = data.get('summary', {}) or {}
            track_info = summary.get('track_info', {}) or {}
            poi_search = summary.get('poi_search', {}) or {}
            poi_results = poi_search.get('results', {}) or {}
            performance = summary.get('performance', {}) or {}
            time_est = summary.get('time_estimate', {})
            
            summary_data = {
                'Categorie': [
                    'ID Analyse',
                    'Fichier source',
                    'Date analyse',
                    'Distance totale (km)',
                    'Nombre de points',
                    'Segments',
                    'Types de POI recherches',
                    'Total POI trouves',
                    'POI nommes',
                    'POI non nommes',
                    'Temps total (s)',
                    'Temps recherche POI (s)',
                    'Temps calcul distances (s)',
                    'Denivele positif (m)',
                    'Denivele negatif (m)',
                    'Altitude min (m)',
                    'Altitude max (m)',
                    'Temps estime (sans pauses)',
                    'Temps estime (avec pauses)',
                    'Temps estime (humain, sans pauses)',
                    'Temps estime (humain, avec pauses)',
                    'Vitesse moyenne',
                    'Allure'
                ],
                'Valeur': [
                    str(data.get('analysis_id', '')),
                    str(data.get('input_file', '')),
                    str(data.get('timestamp', '')),
                    float(track_info.get('total_distance_km', 0.0)),
                    int(track_info.get('total_points', 0)),
                    int(data.get('num_segments', 1)),
                    ', '.join(data.get('poi_types', [])),
                    int(poi_results.get('total_found', 0)),
                    int(poi_results.get('named_count', 0)),
                    int(poi_results.get('unnamed_count', 0)),
                    float(performance.get('total_duration_seconds', 0)),
                    float(performance.get('search_duration_seconds', 0)),
                    float(performance.get('distance_calculation_seconds', 0)),
                    float(track_info.get('total_ascent_m', 0)) ,
                    float(track_info.get('total_descent_m', 0)),
                    float(track_info.get('min_elevation_m', 0)),
                    float(track_info.get('max_elevation_m', 0)),
                    time_est.get('formatted_moving', 'N/A'),
                    time_est.get('formatted_total', 'N/A'),
                    time_est.get('formatted_moving_human', 'N/A'),
                    time_est.get('formatted_total_human', 'N/A'),
                    f"{time_est.get('average_speed_kmh', 0)} km/h",
                    f"{time_est.get('pace_min_per_km', 0)} min/km"
                ]
            }
        
            df = pd.DataFrame(summary_data)
            df.to_excel(writer, sheet_name='Resume', index=False)
        
        except Exception as e:
            logger.error(f"Erreur feuille Resume: {e}")
    
    def _create_poi_sheet(self, writer, data):
        """Cree la feuille POI avec toutes les informations"""
        try:
            pois = data.get('pois', [])
            
            if not pois:
                empty_df = pd.DataFrame({
                    'ID': ['Aucun POI trouve'],
                    'Type': [''],
                    'Nom': [''],
                    'Latitude': [''],
                    'Longitude': [''],
                    'Distance (km)': [''],
                    'Adresse': [''],
                    'Telephone': [''],
                    'Site web': [''],
                    'Horaires': [''],
                    'Description': [''],
                    'Google Maps': [''],
                    'Overpass': ['']
                })
                empty_df.to_excel(writer, sheet_name='POI', index=False)
                return
            
            poi_data = []
            for poi in pois:
                # Recuperation des valeurs de maniere sure
                poi_id = self._get_value(poi, 'id', '')
                poi_type = self._get_value(poi, 'type', '')
                poi_name = self._get_value(poi, 'name', 'Sans nom')
                poi_lat = self._get_float(poi, 'lat', 0.0)
                poi_lon = self._get_float(poi, 'lon', 0.0)
                poi_dist = self._get_float(poi, 'distance', 0.0)
                
                # Formatage de la distance
                if poi_dist > 0:
                    distance_str = f"{poi_dist:.3f}"
                else:
                    distance_str = ''
                
                poi_data.append({
                    'ID': poi_id,
                    'Type': poi_type,
                    'Nom': poi_name,
                    'Latitude': poi_lat,
                    'Longitude': poi_lon,
                    'Distance (km)': distance_str,
                    'Adresse': self._get_value(poi, 'address', ''),
                    'Telephone': self._get_value(poi, 'phone', ''),
                    'Site web': self._get_value(poi, 'website', ''),
                    'Horaires': self._get_value(poi, 'opening_hours', ''),
                    'Description': self._get_value(poi, 'description', ''),
                    'Google Maps': self._get_value(poi, 'google_maps_url', ''),
                    'Overpass': self._get_value(poi, 'overpass_url', '')
                })
            
            df = pd.DataFrame(poi_data)
            df.to_excel(writer, sheet_name='POI', index=False)
            
        except Exception as e:
            logger.error(f"Erreur feuille POI: {e}")
    
    def _create_stats_sheet(self, writer, data):
        """Cree la feuille de statistiques"""
        try:
            summary = data.get('summary', {}) or {}
            poi_search = summary.get('poi_search', {}) or {}
            poi_results = poi_search.get('results', {}) or {}
            
            # Statistiques par type
            by_type = poi_results.get('by_type', {})
            if by_type and isinstance(by_type, dict):
                if len(by_type) > 0:
                    stats_data = {
                        'Type de POI': list(by_type.keys()),
                        'Nombre': list(by_type.values())
                    }
                    df = pd.DataFrame(stats_data)
                    df.to_excel(writer, sheet_name='Statistiques', index=False)
            
            # Distances
            distances = poi_search.get('distances')
            if distances and isinstance(distances, dict):
                min_val = distances.get('min_km')
                max_val = distances.get('max_km')
                avg_val = distances.get('avg_km')
                
                if min_val is not None and max_val is not None and avg_val is not None:
                    dist_data = {
                        'Metrique': ['Distance min (km)', 'Distance max (km)', 'Distance moyenne (km)'],
                        'Valeur': [
                            float(min_val),
                            float(max_val),
                            float(avg_val)
                        ]
                    }
                    df = pd.DataFrame(dist_data)
                    df.to_excel(writer, sheet_name='Distances', index=False)
                
        except Exception as e:
            logger.error(f"Erreur feuille Statistiques: {e}")
    
    def _create_config_sheet(self, writer, data):
        """Cree la feuille de configuration"""
        try:
            radii = data.get('radii', {})
            intervals = data.get('intervals', {})
            
            radii_str = json.dumps(radii, ensure_ascii=False) if radii else '{}'
            intervals_str = json.dumps(intervals, ensure_ascii=False) if intervals else '{}'
            
            config_data = {
                'Parametre': ['Rayons (m)', 'Intervalles (km)'],
                'Valeur': [radii_str, intervals_str]
            }
            df = pd.DataFrame(config_data)
            df.to_excel(writer, sheet_name='Configuration', index=False)
            
        except Exception as e:
            logger.error(f"Erreur feuille Configuration: {e}")
    
    def _apply_styles(self, file_path: Path):
        """Applique un style professionnel au fichier Excel"""
        try:
            wb = load_workbook(file_path)
            
            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
            even_row_fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # Verifier que la feuille a des donnees
                if ws.max_row < 2:
                    continue
                
                # Styler l'en-tete (ligne 1)
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    if cell.value is not None:
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.border = border
                
                # Styler les cellules (lignes 2 a max_row)
                for row_idx in range(2, ws.max_row + 1):
                    for col_idx in range(1, ws.max_column + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        if cell.value is not None:
                            cell.border = border
                            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                            
                            if row_idx % 2 == 0:
                                cell.fill = even_row_fill
                
                # Ajuster la largeur des colonnes
                for col_idx in range(1, ws.max_column + 1):
                    max_length = 0
                    col_letter = get_column_letter(col_idx)
                    for row_idx in range(1, ws.max_row + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        if cell.value is not None:
                            try:
                                cell_len = len(str(cell.value))
                                if cell_len > max_length:
                                    max_length = cell_len
                            except:
                                pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[col_letter].width = adjusted_width
            
            wb.save(file_path)
            
        except Exception as e:
            logger.warning(f"Erreur styles Excel: {e}")