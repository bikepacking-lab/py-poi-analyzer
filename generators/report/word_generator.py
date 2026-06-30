#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from pathlib import Path
from typing import Dict, Any, Optional
import logging
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from ..base_generator import BaseGenerator

logger = logging.getLogger(__name__)

class WordGenerator(BaseGenerator):
    
    def get_format_type(self) -> str:
        return 'docx'
    
    def generate(self, data: Dict[str, Any], output_dir: Path, **kwargs) -> Optional[Path]:
        output_file = output_dir / f'analysis_report_{data.get("analysis_id", "unknown")}.docx'
        
        try:
            doc = Document()
            
            # Titre
            title = doc.add_heading('Rapport d\'Analyse GPS', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Informations generales
            doc.add_heading('Informations generales', level=1)
            doc.add_paragraph(f'ID Analyse: {data.get("analysis_id", "N/A")}')
            doc.add_paragraph(f'Fichier: {data.get("input_file", "N/A")}')
            doc.add_paragraph(f'Date: {data.get("timestamp", "N/A")}')
            
            # Resume
            summary = data.get('summary', {})
            track_info = summary.get('track_info', {})
            poi_search = summary.get('poi_search', {})
            poi_results = poi_search.get('results', {})
            
            performance = summary.get('performance', {})

            doc.add_heading('Resume', level=1)
            doc.add_paragraph(f'Distance totale: {track_info.get("total_distance_km", 0)} km')
            doc.add_paragraph(f'Nombre de points: {track_info.get("total_points", 0)}')
            doc.add_paragraph(f'Total POI: {poi_results.get("total_found", 0)}')
            doc.add_paragraph(f'POI nommes: {poi_results.get("named_count", 0)}')
            doc.add_paragraph(f'POI non nommes: {poi_results.get("unnamed_count", 0)}')

            doc.add_heading('Performance', level=2)
            doc.add_paragraph(f'Temps total: {performance.get("total_duration_seconds", 0):.1f} secondes')
            doc.add_paragraph(f'Temps recherche POI: {performance.get("search_duration_seconds", 0):.1f} secondes')
            doc.add_paragraph(f'Temps calcul distances: {performance.get("distance_calculation_seconds", 0):.1f} secondes')
            
            # TOUS les POI - AUCUNE LIMITE
            doc.add_heading('Liste complete des POI', level=1)
            
            pois = data.get('pois', [])
            
            if not pois:
                doc.add_paragraph('Aucun POI trouve')
            else:
                doc.add_paragraph(f'Total: {len(pois)} POI')
                doc.add_paragraph('')
                
                # Tableau des POI
                table = doc.add_table(rows=1, cols=8)
                table.style = 'Light Grid Accent 1'
                
                # En-tete
                headers = ['#', 'Type', 'Nom', 'Lat', 'Lon', 'Distance', 'Adresse', 'Tel']
                for i, header in enumerate(headers):
                    cell = table.rows[0].cells[i]
                    cell.text = header
                    cell.paragraphs[0].runs[0].bold = True
                
                # TOUS les POI - AUCUNE LIMITE
                for i, poi in enumerate(pois, 1):
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(i)
                    row_cells[1].text = self._get_value(poi, 'type', '')
                    row_cells[2].text = self._get_value(poi, 'name', 'Sans nom')
                    row_cells[3].text = f"{self._get_float(poi, 'lat', 0.0):.5f}"
                    row_cells[4].text = f"{self._get_float(poi, 'lon', 0.0):.5f}"
                    row_cells[5].text = f"{self._get_float(poi, 'distance', 0.0):.2f}"
                    row_cells[6].text = self._get_value(poi, 'address', '')[:30]
                    row_cells[7].text = self._get_value(poi, 'phone', '')
                
                logger.info(f"Word: {len(pois)} POI exportes")
            
            doc.save(str(output_file))
            logger.info(f"Rapport Word genere: {output_file}")
            return output_file
        except Exception as e:
            logger.error(f"Erreur Word: {e}")
            return None
    
    def _get_value(self, obj: Any, key: str, default: str = '') -> str:
        if obj is None:
            return default
        if isinstance(obj, dict):
            val = obj.get(key, default)
            return str(val) if val is not None else default
        if hasattr(obj, key):
            val = getattr(obj, key, default)
            return str(val) if val is not None else default
        return default
    
    def _get_float(self, obj: Any, key: str, default: float = 0.0) -> float:
        if obj is None:
            return default
        try:
            if isinstance(obj, dict):
                val = obj.get(key, default)
            elif hasattr(obj, key):
                val = getattr(obj, key, default)
            else:
                return default
            if val is None:
                return default
            return float(val)
        except (ValueError, TypeError):
            return default